"""Essential PPTX for Data Storage Management System."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0F, 0x2C, 0x3C)
TEAL = RGBColor(0x1A, 0x6B, 0x6B)
ACCENT = RGBColor(0x2A, 0x9D, 0x8F)
LIGHT = RGBColor(0xF4, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x4B, 0x55, 0x63)


def run(p, text, size=16, bold=False, color=DARK):
    p.text = text
    r = p.runs[0]
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def bg(slide, color=LIGHT):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()


def top(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.14))
    s.fill.solid()
    s.fill.fore_color.rgb = TEAL
    s.line.fill.background()


def heading(slide, title):
    top(slide)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.55))
    run(box.text_frame.paragraphs[0], title, size=28, bold=True, color=NAVY)


def footer(slide, n, total):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    run(p, f"Data Storage Management System  |  {n}/{total}", size=11, color=MUTED)
    p.alignment = PP_ALIGN.RIGHT


def bullets(slide, left, top_, width, height, items, size=16):
    box = slide.shapes.add_textbox(left, top_, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run(p, f"• {item}", size=size, color=DARK)
        p.space_after = Pt(10)


def card(slide, left, top_, width, height, title, items):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top_, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(0xC5, 0xD5, 0xD5)
    h = slide.shapes.add_textbox(left + Inches(0.2), top_ + Inches(0.15), width - Inches(0.35), Inches(0.4))
    run(h.text_frame.paragraphs[0], title, size=16, bold=True, color=NAVY)
    bullets(slide, left + Inches(0.2), top_ + Inches(0.55), width - Inches(0.35), height - Inches(0.7), items, size=13)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    total = 7

    # 1 Title
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.9), Inches(13.333), Inches(1.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    t = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.0))
    p = t.text_frame.paragraphs[0]
    run(p, "Data Storage Management System", size=40, bold=True, color=WHITE)
    p.alignment = PP_ALIGN.CENTER
    st = s.shapes.add_textbox(Inches(1.5), Inches(3.4), Inches(10.3), Inches(0.6))
    sp = st.text_frame.paragraphs[0]
    run(sp, "Secure Multi-Tenant Storage with AI Assistance", size=20, color=ACCENT)
    sp.alignment = PP_ALIGN.CENTER
    ft = s.shapes.add_textbox(Inches(1.5), Inches(6.3), Inches(10.3), Inches(0.6))
    fp = ft.text_frame.paragraphs[0]
    run(fp, "Tools  •  Algorithms  •  Features  •  Future Enhancement", size=16, color=WHITE)
    fp.alignment = PP_ALIGN.CENTER

    # 2 Introduction (must)
    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Introduction")
    footer(s, 2, total)
    bullets(
        s,
        Inches(0.7),
        Inches(1.3),
        Inches(12),
        Inches(5.5),
        [
            "Data Storage Management System is a multi-tenant platform for secure file storage and sharing.",
            "It provides User, Admin, and Super Admin portals with JWT authentication and optional 2FA.",
            "Users can upload, organize, share, restore files, and use AI assistance for document summary.",
            "Built with Django REST Framework, React, TypeScript, and PostgreSQL.",
            "Goal: dependable, secure, and scalable storage for smart digital workplaces.",
        ],
        size=18,
    )

    # 3 Tools (must)
    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Most Used Tools")
    footer(s, 3, total)
    card(
        s,
        Inches(0.6),
        Inches(1.3),
        Inches(6.0),
        Inches(5.3),
        "Backend & Database",
        [
            "Python",
            "Django + Django REST Framework",
            "PostgreSQL (SQLite for local test)",
            "SimpleJWT, pyotp, cryptography",
            "Pillow / PyMuPDF for file extraction",
        ],
    )
    card(
        s,
        Inches(6.9),
        Inches(1.3),
        Inches(5.8),
        Inches(5.3),
        "Frontend & Other Tools",
        [
            "React + TypeScript + Vite",
            "Tailwind CSS",
            "PlantUML (CASE / UML diagrams)",
            "Git for version control",
            "Optional: S3/R2, ClamAV, Ollama/Groq",
        ],
    )

    # 4 Algorithms (must)
    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Most Important Algorithms")
    footer(s, 4, total)
    card(s, Inches(0.5), Inches(1.3), Inches(4.0), Inches(2.5), "SHA-256", ["Duplicate file detection", "Share/invite token hashing", "File integrity check"])
    card(s, Inches(4.7), Inches(1.3), Inches(4.0), Inches(2.5), "TOTP (2FA)", ["Time-based one-time password", "Authenticator app login", "Stronger account security"])
    card(s, Inches(8.9), Inches(1.3), Inches(3.9), Inches(2.5), "Malware Scan", ["Heuristic upload scanning", "Block dangerous files", "Optional ClamAV"])
    card(s, Inches(0.5), Inches(4.1), Inches(4.0), Inches(2.5), "TF-IDF", ["Local document summarization", "Sentence ranking", "Privacy-friendly AI"])
    card(s, Inches(4.7), Inches(4.1), Inches(4.0), Inches(2.5), "PBKDF2", ["Password hashing", "Share-link password protection", "No plain-text secrets"])
    card(s, Inches(8.9), Inches(4.1), Inches(3.9), Inches(2.5), "OCR + Intent", ["Read scanned PDFs", "Detect AI request intent", "Select correct file"])

    # 5 Algorithm features (must)
    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Algorithm-Based Features")
    footer(s, 5, total)
    bullets(
        s,
        Inches(0.7),
        Inches(1.3),
        Inches(12),
        Inches(5.5),
        [
            "Secure Upload: malware scan + SHA-256 + quota check before saving files.",
            "Two-Factor Authentication: TOTP codes protect user login.",
            "Safe Sharing: hashed tokens and optional password-protected links.",
            "Duplicate Control: same file content cannot be uploaded twice in one workspace.",
            "AI Assistance: TF-IDF + OCR summarize documents and answer file questions.",
            "Access Control: role-based portals keep User, Admin, and Super Admin separated.",
        ],
        size=18,
    )

    # 6 Future (must) — cloud storage connection emphasized
    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Future Enhancement")
    footer(s, 6, total)
    card(
        s,
        Inches(0.6),
        Inches(1.3),
        Inches(6.0),
        Inches(5.3),
        "Cloud Storage Connection",
        [
            "Full Amazon S3 / Cloudflare R2 integration",
            "Automatic cloud backup of files",
            "Hybrid storage (local + cloud archive)",
            "Cross-region availability",
            "CDN for faster file download",
        ],
    )
    card(
        s,
        Inches(6.9),
        Inches(1.3),
        Inches(5.8),
        Inches(5.3),
        "Other Enhancements",
        [
            "Mobile application support",
            "Advanced AI multi-file analysis",
            "Deeper analytics dashboards",
            "Real-time collaboration alerts",
            "Disaster recovery policies",
        ],
    )

    # 7 Conclusion (must)
    s = prs.slides.add_slide(blank)
    bg(s)
    heading(s, "Conclusion")
    footer(s, 7, total)
    bullets(
        s,
        Inches(0.7),
        Inches(1.4),
        Inches(12),
        Inches(3.5),
        [
            "Data Storage Management System provides secure multi-tenant file management.",
            "Key algorithms improve security, integrity, and AI-based usability.",
            "Modern tools enable scalable web implementation.",
            "Future cloud storage connection will make the system enterprise-ready.",
        ],
        size=18,
    )
    th = s.shapes.add_textbox(Inches(0.7), Inches(5.3), Inches(12), Inches(1.0))
    tp = th.text_frame.paragraphs[0]
    run(tp, "Thank You", size=32, bold=True, color=TEAL)
    tp.alignment = PP_ALIGN.CENTER
    tp2 = th.text_frame.add_paragraph()
    run(tp2, "Questions & Discussion", size=16, color=MUTED)
    tp2.alignment = PP_ALIGN.CENTER

    out = Path(r"d:\Self-Project\Cloud Storage\docs\Data_Storage_Management_System.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(out)
    return out


if __name__ == "__main__":
    build()
